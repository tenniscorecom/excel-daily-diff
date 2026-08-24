"""
src/report.py — 集計結果を Excel に書き出す

集計シートは「行＝積み上げ／延期、列＝種別ごとの日付」の形にする。
列の日付は**ファイル名の日付**（いつ時点の一覧か）であって、案件の日付ではない。
なぜ何件なのかを後から追えるように、明細シートに元の行を残す。
"""

import datetime
import logging
from pathlib import Path

from openpyxl.utils import get_column_letter

from comken.toolbox.excel import Excel, Sheet

from src.diff import DailyDiff
from src.settings import Criteria, Settings, SourceLayout
from src.source import Record

logger = logging.getLogger(__name__)

SUMMARY_SHEET = "集計"
DETAIL_SHEET = "明細"

STATUS_ADDED = "積み上げ"
STATUS_POSTPONED = "延期"

# 集計シートの位置（1始まり）
PLAN_ROW = 1
DATE_ROW = 2
ADDED_ROW = 3
POSTPONED_ROW = 4
LABEL_COL = 1
FIRST_DATE_COL = 2

DATE_NUMBER_FORMAT = "m/d"
COMPARED_DATE_HEADER = "比較日"  # 明細シートで、このツールが付ける列
STATUS_HEADER = "判定"
FROZEN_ROWS = 2  # 集計シートの見出し（種別行・日付行）を固定する

Counts = dict[tuple[str, datetime.date], int]


def write_report(path: Path, diffs: list[DailyDiff], settings: Settings) -> None:
    """集計シートと明細シートを持つブックを作る。

    Args:
        path: 出力先の xlsx パス。
        diffs: 日ごとの差分（古い順）。
        settings: 横軸の期間・種別の並び順・明細シートの見出しに使う列名。
    """
    dates = date_range(settings.start_date, settings.end_date)
    with Excel(path) as excel:
        # 集計シートは表示用のレイアウト（4行・可変列）なので、Excel 側で
        # create_data_sheet ではなく create_sheet を使う。
        # （create_data_sheet だと PY_ 接頭辞が付き、表示用の format / freeze_panes が使えない）
        summary_sheet = excel.create_sheet(SUMMARY_SHEET)
        detail_sheet = excel.create_sheet(DETAIL_SHEET)
        _write_summary(summary_sheet, diffs, settings.criteria, dates)
        _write_detail(detail_sheet, diffs, settings.layout)
        # Excel は with ブロックの正常終了時に自動保存される
    logger.info("出力しました: %s", path)


def date_range(start: datetime.date, end: datetime.date) -> list[datetime.date]:
    """開始日から終了日までの日付を1日ずつ並べて返す。"""
    days = (end - start).days + 1
    return [start + datetime.timedelta(days=i) for i in range(days)]


def detail_headers(layout: SourceLayout) -> list[str]:
    """明細シートの見出し。読み取り元の列名をそのまま使う（config.ini を変えれば追随する）。"""
    return [
        COMPARED_DATE_HEADER,
        STATUS_HEADER,
        layout.key_column,
        layout.date_column,
        layout.plan_column,
        layout.kind_column,
    ]


def _write_summary(
    sheet: Sheet, diffs: list[DailyDiff], criteria: Criteria, dates: list[datetime.date]
) -> None:
    """種別ごとに日付を並べた集計表を書く。"""
    added_counts, postponed_counts = _count_by_plan_and_date(diffs)
    # 比較できなかった日（ファイルが無い日）は 0 と区別できるよう空セルのままにする
    compared_dates = {diff.date for diff in diffs}

    _write_cell(sheet, ADDED_ROW, LABEL_COL, STATUS_ADDED)
    _write_cell(sheet, POSTPONED_ROW, LABEL_COL, STATUS_POSTPONED)

    col = FIRST_DATE_COL
    for plan_prefix in criteria.plan_prefixes:
        # 種別名はグループの先頭列にだけ置く（セルを結合すると並べ替え・集計がしにくい）
        _write_cell(sheet, PLAN_ROW, col, plan_prefix)
        _format(sheet, PLAN_ROW, col, bold=True)
        for date in dates:
            _write_cell(sheet, DATE_ROW, col, date)
            _format(sheet, DATE_ROW, col, number_format=DATE_NUMBER_FORMAT)
            if date in compared_dates:
                _write_cell(sheet, ADDED_ROW, col, added_counts.get((plan_prefix, date), 0))
                _write_cell(sheet, POSTPONED_ROW, col, postponed_counts.get((plan_prefix, date), 0))
            col += 1

    sheet.freeze_panes(_cell_ref(FROZEN_ROWS + 1, LABEL_COL))


def _write_detail(sheet: Sheet, diffs: list[DailyDiff], layout: SourceLayout) -> None:
    """どの行が積み上げ・延期になったかの一覧を書く。"""
    headers = detail_headers(layout)
    rows: list[dict[str, object]] = []
    for diff in diffs:
        rows += [_detail_row(r, diff.date, STATUS_ADDED, layout) for r in _sorted(diff.added)]
        rows += [
            _detail_row(r, diff.date, STATUS_POSTPONED, layout) for r in _sorted(diff.postponed)
        ]
    if not rows:
        # データが無いときは見出しだけ書く
        for column, header in enumerate(headers, start=1):
            _write_cell(sheet, 1, column, header)
        return
    # 表データを二次元配列で write_range に渡す
    matrix = [list(headers), *[[row[header] for header in headers] for row in rows]]
    sheet.write_range(_cell_ref(1, 1) + ":" + _cell_ref(len(matrix), len(headers)), matrix)
    _auto_width(sheet, headers, rows)
    sheet.freeze_panes(_cell_ref(2, 1))


def _count_by_plan_and_date(diffs: list[DailyDiff]) -> tuple[Counts, Counts]:
    """(積み上げ, 延期) の件数を、種別と日付の組ごとに数える。"""
    added: Counts = {}
    postponed: Counts = {}
    for diff in diffs:
        for record in diff.added:
            key = (record.plan_prefix, diff.date)
            added[key] = added.get(key, 0) + 1
        for record in diff.postponed:
            key = (record.plan_prefix, diff.date)
            postponed[key] = postponed.get(key, 0) + 1
    return added, postponed


def _detail_row(
    record: Record, compared_date: datetime.date, status: str, layout: SourceLayout
) -> dict[str, object]:
    return {
        COMPARED_DATE_HEADER: compared_date,
        STATUS_HEADER: status,
        layout.key_column: record.customer_id,
        layout.date_column: record.date,
        layout.plan_column: record.plan,
        layout.kind_column: record.kind,
    }


def _sorted(records: list[Record]) -> list[Record]:
    """日付・種別・キーの順に並べる（毎回同じ並びで出力するため）。"""
    return sorted(records, key=lambda r: (r.date, r.plan_prefix, r.customer_id))


def _cell_ref(row: int, col: int) -> str:
    """(行, 列) の数値を ``A1`` 形式のセル参照に変換する。"""
    return f"{get_column_letter(col)}{row}"


def _write_cell(sheet: Sheet, row: int, col: int, value: object) -> None:
    """Sheet 形式でセル位置と値を、``A1`` 形式を経由して書き込む。"""
    sheet.write_value(_cell_ref(row, col), value)


def _format(sheet: Sheet, row: int, col: int, **kwargs: object) -> None:
    """Sheet 形式でセル書式（太字・表示形式）をまとめて設定する。"""
    sheet.format(_cell_ref(row, col), **kwargs)


def _auto_width(sheet: Sheet, headers: list[str], rows: list[dict[str, object]]) -> None:
    """列ごとに、見出し＋行の最大文字数から列幅を見積もる。

    旧 ``auto_width()`` の代替。日本語（2 幅）やかなは幅が読みにくいので、表示文字数に
    1.2 倍の余裕を持たせた経験的な値で固定する。
    """
    for column, header in enumerate(headers, start=1):
        max_length = len(str(header))
        for row in rows:
            value = row.get(header)
            if value is not None:
                max_length = max(max_length, len(str(value)))
        # 日本語などの全角文字が混ざる可能性に備え、表示幅に余裕を持たせる
        sheet.set_column_width(get_column_letter(column), max(max_length * 1.2, 8))
