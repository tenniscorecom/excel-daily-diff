import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

from src.exceptions import ComparisonFileNotEnoughError
from src.run import _target_files, run
from src.settings import Criteria, FilePattern, Settings, SourceLayout


def _settings(
    folder: Path,
    layout: SourceLayout,
    criteria: Criteria,
    start: datetime.date = datetime.date(2026, 4, 21),
    end: datetime.date = datetime.date(2026, 4, 23),
) -> Settings:
    return Settings(
        folder,
        FilePattern(prefix="一覧_", extension="xlsx"),
        start,
        end,
        folder,
        layout,
        criteria,
    )


def test_target_files_includes_only_previous_and_files_in_range(
    tmp_path: Path, make_book, layout: SourceLayout, criteria: Criteria
) -> None:
    for day in (19, 20, 21, 22, 23, 24):
        make_book(tmp_path / f"一覧_202604{day}.xlsx", [])

    targets = _target_files(_settings(tmp_path, layout, criteria))

    assert [path.name for _, path in targets] == [
        "一覧_20260420.xlsx",
        "一覧_20260421.xlsx",
        "一覧_20260422.xlsx",
        "一覧_20260423.xlsx",
    ]


def test_target_files_uses_previous_when_only_one_file_is_in_range(
    tmp_path: Path, make_book, layout: SourceLayout, criteria: Criteria
) -> None:
    make_book(tmp_path / "一覧_20260420.xlsx", [])
    make_book(tmp_path / "一覧_20260422.xlsx", [])

    targets = _target_files(
        _settings(
            tmp_path,
            layout,
            criteria,
            start=datetime.date(2026, 4, 22),
            end=datetime.date(2026, 4, 22),
        )
    )

    assert [path.name for _, path in targets] == ["一覧_20260420.xlsx", "一覧_20260422.xlsx"]


def test_target_files_accepts_range_starting_at_oldest_file(
    tmp_path: Path, make_book, layout: SourceLayout, criteria: Criteria
) -> None:
    make_book(tmp_path / "一覧_20260421.xlsx", [])
    make_book(tmp_path / "一覧_20260423.xlsx", [])
    make_book(tmp_path / "一覧_20260424.xlsx", [])

    targets = _target_files(_settings(tmp_path, layout, criteria))

    assert [path.name for _, path in targets] == ["一覧_20260421.xlsx", "一覧_20260423.xlsx"]


def test_run_writes_daily_diffs_blank_missing_day_and_compared_date(
    tmp_path: Path, make_book, layout: SourceLayout, criteria: Criteria
) -> None:
    make_book(
        tmp_path / "一覧_20260420.xlsx",
        [("stay", "2026-04-10", "標準A", "完了"), ("drop21", "2026-04-11", "上位B", "予定")],
    )
    make_book(
        tmp_path / "一覧_20260421.xlsx",
        [("stay", "2026-04-10", "標準A", "完了"), ("add21", "2026-04-12", "標準C", "予定")],
    )
    make_book(
        tmp_path / "一覧_20260423.xlsx",
        [("stay", "2026-04-10", "標準A", "完了"), ("add23", "2026-04-13", "上位C", "予定")],
    )

    output_path = run(_settings(tmp_path, layout, criteria))

    workbook = load_workbook(output_path, data_only=True)
    summary = workbook["集計"]
    assert [summary.cell(3, column).value for column in range(2, 8)] == [1, None, 0, 0, None, 1]
    assert [summary.cell(4, column).value for column in range(2, 8)] == [0, None, 1, 1, None, 0]
    assert list(workbook["明細"].values) == [
        ("比較日", "判定", "顧客番号", "予定日", "種別", "状態"),
        (
            datetime.datetime(2026, 4, 21),
            "積み上げ",
            "add21",
            datetime.datetime(2026, 4, 12),
            "標準C",
            "予定",
        ),
        (
            datetime.datetime(2026, 4, 21),
            "延期",
            "drop21",
            datetime.datetime(2026, 4, 11),
            "上位B",
            "予定",
        ),
        (
            datetime.datetime(2026, 4, 23),
            "積み上げ",
            "add23",
            datetime.datetime(2026, 4, 13),
            "上位C",
            "予定",
        ),
        (
            datetime.datetime(2026, 4, 23),
            "延期",
            "add21",
            datetime.datetime(2026, 4, 12),
            "標準C",
            "予定",
        ),
    ]
    workbook.close()


@pytest.mark.parametrize("days", [(21,), (20, 24)])
def test_run_raises_when_comparison_pair_is_unavailable(
    tmp_path: Path, make_book, layout: SourceLayout, criteria: Criteria, days: tuple[int, ...]
) -> None:
    for day in days:
        make_book(tmp_path / f"一覧_202604{day}.xlsx", [])

    with pytest.raises(ComparisonFileNotEnoughError):
        run(_settings(tmp_path, layout, criteria))
